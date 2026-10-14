"""Read the Agenda spreadsheet: the school calendar (🗂 Weeks) and each course tab.

The Agenda is read as a grid of display strings, so the same parser handles the
live Sheets API and an .xlsx snapshot. Columns are found by header text, not by
letter, because the course tabs do not share a layout (Forensics keeps the folder
URL in S, A&P and Biology keep it in R).
"""

from __future__ import annotations

import datetime as dt
import re
from dataclasses import dataclass
from pathlib import Path

from preflight.models import CourseWeek, DayPlan, SchoolWeek
from preflight.sources.gdrive import find_urls

Grid = list[list[str]]

WEEKDAYS = ["Mon", "Tue", "Wed", "Thu", "Fri"]
_MONTHS = {m: i for i, m in enumerate(
    ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"], start=1)}


def week_number(label: str) -> int | None:
    m = re.search(r"\bW(?:k|eek)?\s*0*(\d{1,2})", label or "", re.I)
    return int(m.group(1)) if m else None


def week_label(n: int) -> str:
    return f"Wk {n:02d}"


@dataclass
class AgendaWorkbook:
    tabs: dict[str, Grid]

    def tab(self, name: str) -> Grid:
        if name in self.tabs:
            return self.tabs[name]
        # tolerate emoji/whitespace differences in tab names
        want = _plain(name)
        for key, grid in self.tabs.items():
            if _plain(key) == want:
                return grid
        raise KeyError(f"Agenda tab not found: {name!r}")


def _plain(s: str) -> str:
    return re.sub(r"[^a-z0-9&]+", "", s.lower())


def load_xlsx(path: Path) -> AgendaWorkbook:
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True)
    tabs: dict[str, Grid] = {}
    for ws in wb.worksheets:
        grid: Grid = []
        for row in ws.iter_rows():
            cells = []
            for c in row:
                v = c.value
                if isinstance(v, dt.datetime):
                    v = v.date().isoformat()
                text = "" if v is None else str(v)
                if c.hyperlink is not None and c.hyperlink.target and c.hyperlink.target not in text:
                    text = f"{text} {c.hyperlink.target}".strip()
                cells.append(text)
            grid.append(cells)
        tabs[ws.title] = grid
    return AgendaWorkbook(tabs)


def load_sheets(credentials, spreadsheet_id: str, tab_names: list[str]) -> AgendaWorkbook:
    from googleapiclient.discovery import build

    svc = build("sheets", "v4", credentials=credentials, cache_discovery=False)
    resp = (
        svc.spreadsheets()
        .values()
        .batchGet(
            spreadsheetId=spreadsheet_id,
            ranges=[f"'{t}'" for t in tab_names],
            valueRenderOption="FORMATTED_VALUE",
            dateTimeRenderOption="FORMATTED_STRING",
        )
        .execute()
    )
    tabs: dict[str, Grid] = {}
    for name, vr in zip(tab_names, resp.get("valueRanges", [])):
        tabs[name] = [[str(c) for c in row] for row in vr.get("values", [])]
    return AgendaWorkbook(tabs)


# ---------------------------------------------------------------- calendar

_CLOSURE_RE = re.compile(r"⚠\s*(Mon|Tue|Wed|Thu|Fri)\s+(\d{1,2})/(\d{1,2})\s+(.+)")
_FRIDAY_RE = re.compile(r"Fri\s+\d{1,2}/\d{1,2}\s*=\s*([AB])\s+Friday")


def _parse_iso_or_us(value: str) -> dt.date | None:
    value = value.strip()
    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%Y-%m-%d %H:%M:%S"):
        try:
            return dt.datetime.strptime(value[:19], fmt).date()
        except ValueError:
            continue
    return None


def parse_weeks(grid: Grid) -> dict[int, SchoolWeek]:
    """Parse the 🗂 Weeks engine tab: A=takes over on, C=week label, D=Monday, E=dates & closures."""
    weeks: dict[int, SchoolWeek] = {}
    for row in grid:
        row = row + [""] * (5 - len(row))
        n = week_number(row[2])
        monday = _parse_iso_or_us(row[3]) or _parse_iso_or_us(row[0])
        if n is None or monday is None:
            continue
        raw = row[4]
        closures: dict[dt.date, str] = {}
        for m in _CLOSURE_RE.finditer(raw):
            day = dt.date(monday.year, int(m.group(2)), int(m.group(3)))
            if day < monday - dt.timedelta(days=200):  # crosses New Year
                day = day.replace(year=monday.year + 1)
            reason = m.group(4).strip()
            if re.search(r"closed|holiday|in-service|labor day|memorial|break|election|thanksgiving",
                         reason, re.I):
                closures[day] = reason
        fm = _FRIDAY_RE.search(raw)
        weeks[n] = SchoolWeek(label=week_label(n), monday=monday,
                              friday_side=fm.group(1) if fm else "", closures=closures, raw=raw)
    return weeks


def week_for_date(weeks: dict[int, SchoolWeek], day: dt.date) -> int | None:
    for n, w in weeks.items():
        if w.monday <= day <= w.monday + dt.timedelta(days=6):
            return n
    return None


# ---------------------------------------------------------------- course tabs

def _header_index(grid: Grid) -> tuple[int, dict[str, int]]:
    for i, row in enumerate(grid[:15]):
        if row and row[0].strip() == "Wk":
            return i, {h.strip(): j for j, h in enumerate(row) if h.strip()}
    raise ValueError("course tab has no header row starting with 'Wk'")


def _col(headers: dict[str, int], *needles: str) -> int | None:
    for name, j in headers.items():
        low = name.lower()
        if all(n.lower() in low for n in needles):
            return j
    return None


def parse_course_week(grid: Grid, course: str, week_no: int, school_week: SchoolWeek) -> CourseWeek | None:
    head_row, headers = _header_index(grid)
    target = None
    for row in grid[head_row + 1:]:
        if row and week_number(row[0]) == week_no and row[0].strip().lower().startswith("w"):
            target = row
            break
    if target is None:
        return None
    width = max(headers.values()) + 1
    row = target + [""] * (width - len(target))

    def cell(*needles: str) -> str:
        j = _col(headers, *needles)
        return row[j].strip() if j is not None and j < len(row) else ""

    days: list[DayPlan] = []
    for i, wd in enumerate(WEEKDAYS):
        plan_j = _col(headers, wd, "plan")
        if plan_j is None:
            continue
        note_j = _col(headers, wd, "note")
        days.append(DayPlan(
            date=school_week.monday + dt.timedelta(days=i),
            weekday=wd,
            plan=row[plan_j].strip(),
            note=row[note_j].strip() if note_j is not None else "",
        ))

    named = {
        "student_url": cell("_studenturl"),
        "teacher_url": cell("_teacherurl"),
        "folder_url": cell("_folderurl"),
        "extra_url": cell("_extraurl"),
        "turn_in_url": cell("_turninurl"),
    }
    # Some rows hold a URL in the display column but not in the hidden one.
    for key, display in (("folder_url", "folder"), ("turn_in_url", "turn it in")):
        if not named[key]:
            urls = find_urls(cell(display))
            named[key] = urls[0] if urls else ""

    known = {u for u in named.values() if u}
    other: list[str] = []
    for value in row:
        for u in find_urls(value):
            if u not in known and u not in other:
                other.append(u)

    return CourseWeek(
        course=course,
        week_label=week_label(week_no),
        topic=cell("topic"),
        dates_cell=cell("dates"),
        days=days,
        todo=cell("to-do"),
        lab_prep=cell("lab"),
        print_set=cell("print set"),
        other_urls=other,
        row_cells={h: (row[j] if j < len(row) else "") for h, j in headers.items()},
        **named,
    )


# ---------------------------------------------------------------- side tabs

def worksheets_for_week(grid: Grid, course: str, week_no: int) -> list[dict[str, str]]:
    """Rows of the 📝 Worksheets tab for this course and week. Columns: Course, Wk, Topic, Document, Type, Link, _fileID."""
    out = []
    for row in grid[1:]:
        row = row + [""] * (7 - len(row))
        if _plain(row[0]) != _plain(course):
            continue
        if week_number(row[1]) != week_no:
            continue
        out.append({"course": row[0], "wk": row[1], "topic": row[2], "document": row[3],
                    "type": row[4], "link": row[5], "file_id": row[6].strip()})
    return out


def turnin_for_week(grid: Grid, course: str, week_no: int) -> list[dict[str, str]]:
    """Rows of 🗂 TurnIn Map: assignment key, target folder id, course, week."""
    out = []
    for row in grid[1:]:
        row = row + [""] * (4 - len(row))
        if _plain(course) in _plain(row[2]) and week_number(row[3]) == week_no:
            out.append({"key": row[0], "folder_id": row[1].strip(), "course": row[2], "week": row[3]})
    return out


def parse_month_day(text: str, year: int) -> list[dt.date]:
    """All M/D dates in free text, e.g. '10/19', 'Tue 10/13'."""
    out = []
    for m in re.finditer(r"(?<![\d/])(\d{1,2})/(\d{1,2})(?![\d/])", text or ""):
        mo, d = int(m.group(1)), int(m.group(2))
        if 1 <= mo <= 12 and 1 <= d <= 31:
            try:
                out.append(dt.date(year if mo >= 7 else year + 1, mo, d))
            except ValueError:
                pass
    return out


def parse_month_name_dates(text: str, school_year_start: int) -> list[dt.date]:
    out = []
    for m in re.finditer(r"\b(Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[a-z]*\.?\s+(\d{1,2})\b", text or ""):
        mo = _MONTHS[m.group(1)]
        try:
            out.append(dt.date(school_year_start if mo >= 7 else school_year_start + 1, mo, int(m.group(2))))
        except ValueError:
            pass
    return out
