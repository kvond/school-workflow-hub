"""Build small synthetic Agendas, decks and Drive snapshots for tests."""

from __future__ import annotations

import datetime as dt
import io
import json
from pathlib import Path

import openpyxl
from pptx import Presentation
from pptx.util import Inches

from preflight.config import Config, CourseConfig, Pacing

WEEKS_TAB = "🗂 Weeks"
COURSE_TAB = "🔬 Forensics"
HEADERS = ["Wk", "Dates & closures", "Topic", "Mon — plan", "📝 Mon note", "Wed — plan", "📝 Wed note",
           "Fri — plan", "📝 Fri note", "✅ To-do", "🧪 Lab / prep", "🎞 Student deck", "👩‍🏫 Teacher deck",
           "📁 Folder", "🎞 2nd deck", "_studentURL", "_teacherURL", "", "_folderURL", "_extraURL",
           "📤 Turn it in", "_turnInURL"]

DECK_ID = "D" * 33
FOLDER_ID = "F" * 33
FORM_ID = "R" * 44
ROOT_ID = "ROOT" + "r" * 29
TURNIN_ID = "T" * 33


def make_config(**over) -> Config:
    course = CourseConfig(name="Forensics", tab=COURSE_TAB, meets=["Mon", "Wed"], friday_side="A",
                          marker_prefix="F", students=24, groups=6, drive_root=ROOT_ID,
                          sections=over.pop("sections", []))
    return Config(agenda_id="x", school_year_start=2026, courses=[course], pacing=Pacing(), **over)


def make_agenda(path: Path, row: dict[str, str], weeks_e: str | None = None, extra_rows=()) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = WEEKS_TAB
    ws.append(["ENGINE"])
    ws.append(["takes over on", "course row", "week", "Monday", "dates & closures"])
    ws.append([dt.datetime(2026, 10, 12), 10, "Wk 07", dt.datetime(2026, 10, 12),
               "Oct 12 – Oct 16, 2026\n⚠ Fri 10/16 In-Service — Schools Closed\nFri 10/16 = B Friday → A&P"])
    ws.append([dt.datetime(2026, 10, 19), 11, "Wk 08", dt.datetime(2026, 10, 19),
               weeks_e or "Oct 19 – Oct 23, 2026\nFri 10/23 = A Friday → Bio + Forensics"])
    ws.append([dt.datetime(2026, 10, 26), 12, "Wk 09", dt.datetime(2026, 10, 26),
               "Oct 26 – Oct 30, 2026\nFri 10/30 = B Friday → A&P"])
    c = wb.create_sheet(COURSE_TAB)
    c.append(["🔬 Forensics — 2026-27"])
    c.append([""])
    c.append(HEADERS)
    for r in [row, *extra_rows]:
        c.append([r.get(h, "") for h in HEADERS])
    w = wb.create_sheet("📝 Worksheets")
    w.append(["Course", "Wk", "Topic", "Document", "Type", "Link", "_fileID"])
    t = wb.create_sheet("🗂 TurnIn Map")
    t.append(["assignment key", "target folder id", "course", "week"])
    t.append(["🔬 Forensics · Wk 08 — History of Forensics", TURNIN_ID, "🔬 Forensics", "Wk 08"])
    wb.save(path)


def week8_row(**over) -> dict[str, str]:
    row = {
        "Wk": "Wk 08",
        "Topic": "History of Forensics",
        "Mon — plan": "10/19\nHistory of Forensics webquest",
        "Wed — plan": "10/21\nTimeline build",
        "Fri — plan": "10/23\nSecond vote and wrap-up",
        "🧪 Lab / prep": "BUY: Carrots, sugar · PULL: Trays · ✔ on order — yeast, balloons; 🛒 tape",
        "📁 Folder": "Week 08 ↗",
        "_studentURL": f"https://docs.google.com/presentation/d/{DECK_ID}/copy",
        "_teacherURL": f"https://docs.google.com/presentation/d/{DECK_ID}/edit",
        "_folderURL": f"https://drive.google.com/drive/folders/{FOLDER_ID}",
        "_turnInURL": f"https://docs.google.com/forms/d/{FORM_ID}/viewform",
    }
    row.update(over)
    return row


def make_deck(slides: list[str | tuple[str, str]]) -> bytes:
    """Each slide is text, or (text, speaker notes). One text box per paragraph block."""
    prs = Presentation()
    blank = prs.slide_layouts[6]
    for item in slides:
        text, notes = (item if isinstance(item, tuple) else (item, ""))
        s = prs.slides.add_slide(blank)
        box = s.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(9), Inches(6))
        box.text_frame.text = text
        if notes:
            s.notes_slide.notes_text_frame.text = notes
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


GOOD_DECK = [
    "[[MARKER-INVENTORY:F08]] NOTES=6 DRAFT=1 OPTIONAL=0 BANK_TERMS=2 "
    "NOTES_IDS=CA1-Q,CA1-RULE,CA1-PB,CA1-MASTERY,X1,X2 DRAFT_IDS=CA1-MASTERY OPTIONAL_IDS=-",
    "TEACHER NOTE — do not project\nCritical aspect: contact leaves a trace\nWhat varies: contact. Held constant: the scene.",
    "Critical aspect: contact leaves a trace\nTwo rooms, one touched. Which one holds evidence? [[NOTES:F08:CA1-Q]]",
    "Critical aspect: contact leaves a trace\nCASE A — gloves on\nCASE B — gloves off",
    "Critical aspect: contact leaves a trace\nFinish this sentence as a rule: when two things touch, ______. [[NOTES:F08:CA1-RULE]]",
    "Pattern break / Critical aspect: contact leaves a trace\nA scene with no trace at all. Why? [[NOTES:F08:CA1-PB]]",
    "Critical aspect: contact leaves a trace\nGetting Started\nName the trace.\nWorking On It\nMastery",
    ("Your answer — Critical Aspect 1 · contact leaves a trace\nYour first answer [[NOTES:F08:CA1-MASTERY]]\n"
     "Your revised answer [[DRAFT:F08:CA1-MASTERY]]", "REVISION PROMPT — copy everything between the lines"),
    "Extra one [[NOTES:F08:X1]]\nExtra two [[NOTES:F08:X2]]",
    "CONCEPT BANK\nDefine these in your own words\nThen explain how two of them connect.\n"
    "Trace [[BANK:F08:Trace]]\nContact [[BANK:F08:Contact]]",
]


def make_snapshot(root: Path, deck_slides=None, files_extra=(), row=None, weeks_e=None, extra_rows=(),
                  schoology=None, folder=True) -> Path:
    root.mkdir(parents=True, exist_ok=True)
    make_agenda(root / "agenda.xlsx", row or week8_row(), weeks_e, extra_rows)
    drive = root / "drive"
    (drive / "exports").mkdir(parents=True, exist_ok=True)
    files = [
        {"id": ROOT_ID, "name": "Forensics", "mimeType": "application/vnd.google-apps.folder", "parents": []},
        {"id": DECK_ID, "name": "Week 08 — History of Forensics (VT)",
         "mimeType": "application/vnd.google-apps.presentation", "parents": [FOLDER_ID]},
        {"id": TURNIN_ID, "name": "🔬 Forensics", "mimeType": "application/vnd.google-apps.folder", "parents": []},
        {"id": FORM_ID, "name": "Forensics — Wk 08 — Turn In", "mimeType": "application/vnd.google-apps.form",
         "parents": []},
    ]
    if folder:
        files.append({"id": FOLDER_ID, "name": "Week 08 - History of Forensics",
                      "mimeType": "application/vnd.google-apps.folder", "parents": [ROOT_ID]})
    files.extend(files_extra)
    (drive / "files.json").write_text(json.dumps(files))
    (drive / "exports" / f"{DECK_ID}.pptx").write_bytes(make_deck(deck_slides or GOOD_DECK))
    (root / "schoology.json").write_text(json.dumps(schoology or {}))
    return root
